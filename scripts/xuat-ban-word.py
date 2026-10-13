# -*- coding: utf-8 -*-
"""Dung ban Word cua bao cao theo dung quy dinh trinh bay cua Khoa KT-CN DH Tra Vinh."""
import re, os, sys, struct
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK,
                            WD_TAB_ALIGNMENT, WD_TAB_LEADER)
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = '/home/user/Project-Final-TVU'
DOCS = os.path.join(ROOT, 'docs', 'bao-cao')
IMGS = '/home/user/Project-Final-TVU/docs/images/bao-cao'

# ── kich thuoc PNG doc truc tiep tu IHDR, khong can thu vien anh ──────────────
def png_size(path):
    with open(path, 'rb') as f:
        head = f.read(26)
    if head[:8] != b'\x89PNG\r\n\x1a\n':
        return None
    w, h = struct.unpack('>II', head[16:24])
    return w, h

def figure_files():
    out = {}
    for f in os.listdir(IMGS):
        m = re.match(r'hinh-([0-9a-z]+)-([0-9]+)-', f)
        if m:
            out[f"{m.group(1).upper()}.{m.group(2)}"] = os.path.join(IMGS, f)
    return out

FIGS = figure_files()

CAPTIONS = {}
# Cac dong tu dung de dan mot hinh trong phan chu. Ban dau chi bat 'thể hiện'
# va chi o DAU dong, nen 10 trong 31 hinh khong lay duoc mo ta — khong phai vi
# chung khong duoc nhac toi, ma vi cau dan viet bang dong tu khac ('minh hoạ',
# 'mô tả', 'so sánh') hoac nam giua cau. Mot hinh khong co mo ta thi chu thich
# chi con 'Hình 3.4 (Nguồn: tác giả)', va danh muc hinh o phan dau cung trong.
DONG_TU_DAN = r'(?:thể hiện|minh hoạ|minh họa|mô tả|so sánh|trình bày)'

def load_captions():
    """Lay mo ta hinh tu cau dan 'Hình X.Y <dong tu> ...' trong chinh phan chu."""
    mau = re.compile(r'Hình ([0-9A-Z]+\.[0-9]+)\s+' + DONG_TU_DAN + r'\s+')
    for f in sorted(os.listdir(DOCS)):
        if not f.endswith('.md'):
            continue
        noi_dung = open(os.path.join(DOCS, f), encoding='utf-8').read()
        # Gop cac dong cua cung mot doan lai truoc khi doi chieu. Markdown ngat
        # dong o cot 80 nen mot cau dan nhu "Hình 3.4 thể hiện" thuong bi cat
        # ngay giua: dong tu o cuoi dong nay, mo ta o dau dong sau.
        for doan in re.split(r'\n\s*\n', noi_dung):
            doan = ' '.join(doan.split())
            if doan.startswith('[Hình'):
                continue
            # Quet TUNG vi tri cau dan roi cat tay, khong dung mot nhom bat
            # tham lam: mot doan hay co hai cau dan lien nhau ("Hình 3.17 thể
            # hiện ... Hình 3.18 thể hiện ..."), va nhom tham lam se nuot cau
            # thu hai vao mo ta cua cau thu nhat — hinh 3.18 mat mo ta.
            moc = list(mau.finditer(doan))
            for k, m in enumerate(moc):
                so = m.group(1)
                if so in CAPTIONS:
                    continue
                het = moc[k + 1].start() if k + 1 < len(moc) else len(doan)
                t = doan[m.end():het].strip()
                # Cat o cho cau chuyen y hoac o het cau dau tien.
                t = re.split(r'\s*,?\s*(?:trong đó|còn Hình|và Hình)\b', t)[0]
                t = re.split(r'(?<=[^0-9])\.\s', t)[0]
                t = t.rstrip(' .,;:').replace('**', '').replace('`', '').strip()
                if t:
                    CAPTIONS[so] = t[0].upper() + t[1:]
load_captions()

# ── dinh dang chung ──────────────────────────────────────────────────────────
FONT = 'Times New Roman'

def set_font(run, size=13, bold=False, italic=False, mono=False):
    name = 'Consolas' if mono else FONT
    run.font.name = name          # python-docx dat rFonts dung vi tri trong rPr
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rFonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for a in ('w:cs', 'w:eastAsia'):
        rFonts.set(qn(a), name)

def body_format(p, size=13, spacing=1.5, before=6, after=6,
                align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.alignment = align

INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)')

HTML_TAG = re.compile(r'<(/?)([a-zA-Z][a-zA-Z0-9]*)[^>]*>')
_da_canh_bao = set()


def lam_sach_html(text):
    """Bo the HTML con sot trong Markdown.

    Word khong hieu HTML, nen mot the <br> viet trong nguon se in ra nguyen
    van "<br>" giua trang giay. Da dinh mot lan o bang ky ten cua de cuong.
    <br> doi thanh dau xuong dong; the la thi bo di kem mot dong canh bao,
    vi im lang la cach loi nay quay lai."""
    text = re.sub(r'<br\s*/?>', '\n', text, flags=re.I)
    for m in HTML_TAG.finditer(text):
        ten = m.group(2).lower()
        if ten not in _da_canh_bao:
            _da_canh_bao.add(ten)
            print('  canh bao: bo the HTML <%s> — Word khong hieu HTML' % ten)
    return HTML_TAG.sub('', text)


def add_inline(p, text, size=13, base_bold=False):
    text = lam_sach_html(text)
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)   # bo lien ket
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            set_font(p.add_run(part[2:-2]), size, bold=True)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            set_font(p.add_run(part[1:-1]), size, italic=True)
        elif part.startswith('`') and part.endswith('`'):
            set_font(p.add_run(part[1:-1]), size - 1.5, mono=True)
        else:
            for i, dong in enumerate(part.split('\n')):
                if i:
                    p.add_run().add_break()
                if dong:
                    set_font(p.add_run(dong), size, bold=base_bold)

# ── danh so: moi danh sach mot day rieng ─────────────────────────────────────
# Style 'List Number' cua Word la MOT day so duy nhat cho ca tai lieu. Dung
# nguyen no thi muc "Dong gop du kien" noi tiep so cua muc "Nhiem vu" truoc do
# (1..9 roi 10, 11, 12), va danh muc tai lieu tham khao mo dau bang "13." —
# trong nhu bi thieu mat 12 muc dau. Moi danh sach vi the phai co numId rieng.
_DANH_SO = {'abstract': None}

def _phan_numbering(doc):
    return doc.part.numbering_part.element

def _tao_abstract_num(doc):
    """Dinh nghia mot kieu danh so thap phan ba cap, dung chung cho moi danh sach."""
    goc = _phan_numbering(doc)
    da_co = [int(a.get(qn('w:abstractNumId')))
             for a in goc.findall(qn('w:abstractNum'))]
    aid = max(da_co) + 1 if da_co else 0
    a = OxmlElement('w:abstractNum')
    a.set(qn('w:abstractNumId'), str(aid))
    for cap in range(3):
        lvl = OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'), str(cap))
        for ten, val in (('w:start', '1'), ('w:numFmt', 'decimal'),
                         ('w:lvlText', '%%%d.' % (cap + 1)), ('w:lvlJc', 'left')):
            e = OxmlElement(ten); e.set(qn('w:val'), val); lvl.append(e)
        pPr = OxmlElement('w:pPr'); ind = OxmlElement('w:ind')
        ind.set(qn('w:left'), str(360 + 360 * cap)); ind.set(qn('w:hanging'), '360')
        pPr.append(ind); lvl.append(pPr)
        a.append(lvl)
    # Trong numbering.xml moi the w:abstractNum phai dung TRUOC moi the w:num.
    dau = goc.find(qn('w:num'))
    dau.addprevious(a) if dau is not None else goc.append(a)
    return aid

def day_so_moi(doc):
    """Cap mot numId chua dung, tuc mot bo dem bat dau lai tu 1."""
    if _DANH_SO['abstract'] is None:
        _DANH_SO['abstract'] = _tao_abstract_num(doc)
    goc = _phan_numbering(doc)
    da_co = [int(n.get(qn('w:numId'))) for n in goc.findall(qn('w:num'))]
    nid = max(da_co) + 1 if da_co else 1
    n = OxmlElement('w:num'); n.set(qn('w:numId'), str(nid))
    ab = OxmlElement('w:abstractNumId'); ab.set(qn('w:val'), str(_DANH_SO['abstract']))
    n.append(ab)
    # Bo dem nam o w:abstractNum, KHONG phai o w:num. Nhieu numId cung tro ve
    # mot abstractNum thi van dem tiep cung mot day — cap numId moi thoi la
    # chua du. Moi day vi the phai tu khai lai diem bat dau bang startOverride.
    for cap in range(3):
        ov = OxmlElement('w:lvlOverride'); ov.set(qn('w:ilvl'), str(cap))
        st = OxmlElement('w:startOverride'); st.set(qn('w:val'), '1')
        ov.append(st); n.append(ov)
    goc.append(n)
    return nid

def gan_day_so(p, nid, cap=0):
    """Gan doan van vao day so nid. Dung get_or_add_* de the w:numPr roi
    dung VI TRI ma lich do OOXML quy dinh trong w:pPr — chen sai cho thi
    LibreOffice tu choi mo tep, khong bao loi gi ro rang."""
    numPr = p._element.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = cap
    numPr.get_or_add_numId().val = nid

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)

def add_figure(doc, num, width_cm):
    path = FIGS.get(num)
    if not path:
        p = doc.add_paragraph(); body_format(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_font(p.add_run(f'[ CHUA CO HINH {num} ]'), 13, bold=True)
        return
    size = png_size(path)
    w_cm = width_cm
    if size:
        w_px, h_px = size
        # Hai rang buoc nguoc chieu nhau, va chieu DOC DUOC thang.
        #
        # Tran chieu cao giu cho hinh khong chiem tron trang. Nhung ap tran do
        # len mot hinh cao-hep thi no bop be ngang lai theo: bieu do use case
        # 1320x2415 o tran 7cm chi con 3,8cm ngang, so do quy trinh bay buoc
        # con 3,2cm — in ra khong doc duoc chu nao. Mot so do khong doc duoc
        # thi khong con la so do.
        #
        # Nen them SAN be ngang 8,5cm: hinh nao bi tran chieu cao bop hep hon
        # muc do thi duoc cao hon tran, toi 18cm. Hinh cao hon nua moi bi thu
        # lai theo chieu cao.
        w_cm = min(w_cm, 15.0)
        cao = h_px / w_px * w_cm
        if cao > 7:
            w_cm = 7 * w_px / h_px
            if w_cm < 8.5:
                w_cm = 8.5
                if h_px / w_px * w_cm > 18:
                    w_cm = 18 * w_px / h_px
    p = doc.add_paragraph(); body_format(p, before=6, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.add_run().add_picture(path, width=Cm(w_cm))
    cap = doc.add_paragraph()
    body_format(cap, size=12, spacing=1.0, before=2, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    desc = CAPTIONS.get(num, '')
    nguon = 'Nguồn: tác giả'
    set_font(cap.add_run(f'Hình {num}'), 12, bold=True)
    set_font(cap.add_run(f' — {desc} ({nguon})' if desc else f' ({nguon})'), 12, italic=True)

def add_muc_luc(doc, rows):
    """Muc luc kieu tai lieu in: ten muc — dau cham noi — so trang.

    Ve bang mot doan van co mot diem dung tab CAN PHAI kem dau dan la dau
    cham. Dung bang hai cot thi so trang khong thang hang voi nhau khi ten
    muc dai ngan khac nhau, va nhin ra ngay la bang chu khong phai muc luc."""
    p = doc.add_paragraph()
    body_format(p, size=14, spacing=1.15, before=6, after=10,
                align=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(p.add_run('MỤC LỤC'), 14, bold=True)
    for ten, trang in rows[1:]:          # bo dong tieu de 'Muc | Trang'
        tho = ten.replace('**', '').strip()
        # Muc cap hai ('2.1.', '6.4.') thut vao; muc cap mot in dam.
        cap_hai = bool(re.match(r'^\d+\.\d+\.', tho))
        p = doc.add_paragraph()
        body_format(p, size=13, spacing=1.15, before=1, after=1,
                    align=WD_ALIGN_PARAGRAPH.LEFT)
        if cap_hai:
            p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(15.8), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        set_font(p.add_run(tho), 13, bold=not cap_hai)
        set_font(p.add_run('\t' + trang.strip()), 13, bold=not cap_hai)

def doc_can_le(o):
    """':---:' -> can giua, '---:' -> can phai, con lai -> can trai."""
    o = o.strip()
    if o.startswith(':') and o.endswith(':'): return WD_ALIGN_PARAGRAPH.CENTER
    if o.endswith(':'):                       return WD_ALIGN_PARAGRAPH.RIGHT
    return WD_ALIGN_PARAGRAPH.LEFT

def add_table(doc, rows, width_cm, can_le=None):
    # Bang co dong tieu de dung hai chu 'Muc' va 'Trang' la MUC LUC, khong
    # phai mot bang du lieu — ve bang dau cham noi cho dung kieu tai lieu in.
    if len(rows) > 1 and [c.strip() for c in rows[0][:2]] == ['Mục', 'Trang']:
        add_muc_luc(doc, rows)
        return
    ncol = max(len(r) for r in rows)
    rows = [r + [''] * (ncol - len(r)) for r in rows]
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ''
            p = cell.paragraphs[0]
            cot = (can_le[j] if can_le and j < len(can_le)
                   else WD_ALIGN_PARAGRAPH.LEFT)
            body_format(p, size=11, spacing=1.0, before=1, after=1, align=cot)
            add_inline(p, cell_text, size=11, base_bold=(i == 0))
            if i == 0:
                shade(cell, 'EFEDE7')
    # Bang dai vo sang trang sau thi trang do mat dong tieu de, va nguoi doc
    # khong con biet cot nao la cot nao. Danh dau dong dau la dong tieu de de
    # Word tu lap lai no o moi trang.
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader')
    th.set(qn('w:val'), 'true')
    trPr.append(th)
    # Dong hoan toan rong giu mot chieu cao toi thieu: do la cach mot bang ky
    # ten chua duoc cho de ky. Khong dat thi dong rong cao bang mot dong chu.
    for i, row in enumerate(rows):
        if i and not any(c.strip() for c in row):
            pr = t.rows[i]._tr.get_or_add_trPr()
            h = OxmlElement('w:trHeight')
            h.set(qn('w:val'), '1400')     # ~2,5cm — vua mot cho ky ten
            h.set(qn('w:hRule'), 'atLeast')
            pr.append(h)
    return t

def add_code(doc, lines):
    p = doc.add_paragraph()
    body_format(p, size=10.5, spacing=1.0, before=6, after=6, align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.left_indent = Cm(0.6)
    for i, l in enumerate(lines):
        if i: p.add_run().add_break()
        set_font(p.add_run(l), 10.5, mono=True)

# Dong ke phan cach cua bang PHAI co it nhat mot dau gach. Mau cu chi doi
# "toan ky tu gach, hai cham, khoang trang va gach dung", nen mot dong RONG
# kieu "|  |  |" cung khop va bi bo di im lang — dung ba dong trong chua cho
# ky ten cuoi de cuong bien mat ma khong bao gi.
TABLE_SEP = re.compile(r'^\s*\|?[\s:\-|]*-[\s:\-|]*\|[\s:\-|]*$')
# Dong toan dau cham trong Markdown: dong ke de nguoi doc viet tay.
DONG_CHAM = re.compile(r'^\.{10,}$')

def render_markdown(doc, path, width_cm, first_chapter_break=True):
    lines = open(path, encoding='utf-8').read().split('\n')
    i, para, first_h1 = 0, [], True
    # numId cua danh sach so dang mo; None nghia la doan truoc khong phai
    # muc danh so, nen muc tiep theo mo mot day moi bat dau tu 1.
    day_so_dang_mo = None
    def flush():
        nonlocal para
        if para:
            p = doc.add_paragraph(); body_format(p)
            add_inline(p, ' '.join(para))
            para = []
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            # Dong trong KHONG dong day so: mot danh sach "thua" (co dong trong
            # giua cac muc) van la mot danh sach.
            flush(); i += 1; continue
        if not re.match(r'^\d+\.\s+', s):
            # Bat cu khoi nao khac — tieu de, bang, gach dau dong, doan van —
            # deu ket thuc day so dang mo. Dong noi tiep cua mot muc danh sach
            # khong toi duoc day: nhanh danh sach o duoi da nuot chung.
            day_so_dang_mo = None
        m = re.match(r'^(#{1,4})\s+(.*)', s)
        if m:
            flush()
            lvl, txt = len(m.group(1)), m.group(2).strip()
            p = doc.add_paragraph()
            if lvl == 1:
                if not first_h1 or first_chapter_break:
                    p.paragraph_format.page_break_before = True
                first_h1 = False
                body_format(p, size=14, spacing=1.5, before=0, after=12,
                            align=WD_ALIGN_PARAGRAPH.CENTER)
                set_font(p.add_run(txt.upper()), 14, bold=True)
            else:
                sizes = {2: 13.5, 3: 13, 4: 13}
                body_format(p, size=sizes[lvl], spacing=1.5, before=12, after=6,
                            align=WD_ALIGN_PARAGRAPH.LEFT)
                set_font(p.add_run(txt), sizes[lvl], bold=True)
                p.paragraph_format.keep_with_next = True
            i += 1; continue
        if s.startswith('```'):
            flush(); i += 1; buf = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                buf.append(lines[i]); i += 1
            i += 1
            add_code(doc, buf); continue
        if re.match(r'^\[Hình ([0-9A-Z]+\.[0-9]+)\]$', s):
            flush()
            add_figure(doc, re.match(r'^\[Hình ([0-9A-Z]+\.[0-9]+)\]$', s).group(1), width_cm)
            i += 1; continue
        # Chu thich bang, dat NGAY TRUOC bang. Quy dinh doi "danh so bang
        # bieu, hinh ve, do thi"; truoc day chi hinh duoc danh so con bang thi
        # khong, nen moi bang trong quyen la mot bang khong ten.
        #
        # Dat truoc bang chu khong phai sau, giong bao cao mau cua khoa: nguoi
        # doc biet minh sap doc bang gi truoc khi doc so lieu.
        m_bang = re.match(r'^\[Bảng ([0-9A-Z]+\.[0-9]+)\]\s*(.*)$', s)
        if m_bang:
            flush()
            cap = doc.add_paragraph()
            body_format(cap, size=12, spacing=1.0, before=8, after=2,
                        align=WD_ALIGN_PARAGRAPH.CENTER)
            set_font(cap.add_run(f'Bảng {m_bang.group(1)}'), 12, bold=True)
            mo_ta = m_bang.group(2).strip().rstrip('.')
            set_font(cap.add_run(f' — {mo_ta} (Nguồn: tác giả)' if mo_ta
                                 else ' (Nguồn: tác giả)'), 12, italic=True)
            i += 1; continue
        if s.startswith('|'):
            flush(); rows, can_le = [], []
            while i < len(lines) and lines[i].strip().startswith('|'):
                raw = lines[i].strip()
                if TABLE_SEP.match(raw):
                    # Dong phan cach khong phai du lieu, nhung no cho biet
                    # tung cot can le the nao.
                    can_le = [doc_can_le(c) for c in raw.strip().strip('|').split('|')]
                else:
                    cells = [c.strip() for c in raw.strip('|').split('|')]
                    rows.append(cells)
                i += 1
            if rows: add_table(doc, rows, width_cm, can_le)
            continue
        if DONG_CHAM.match(s):
            # Dong ke trong de nguoi cham viet tay. Mot chuoi dau cham la MOT
            # "tu" lien khoi: can deu chi keo gian khoang trang GIUA cac tu,
            # ma o day khong co khoang trang nao, nen dong dung lai o dau het
            # cham — do duoc 243,7pt tren 453,5pt be chu, tuc hon nua dong.
            # Thay bang mot tab co dan dau la dau cham, diem dung dat sat le
            # phai: Word tu rai cham kin be ngang, khong phu thuoc so cham go
            # trong Markdown hay co chu dang dung.
            flush()
            p = doc.add_paragraph()
            body_format(p)
            p.paragraph_format.tab_stops.add_tab_stop(
                Cm(width_cm), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            set_font(p.add_run('\t'), 13)
            i += 1; continue
        if s.startswith('\\khoiky '):
            # Khoi ky ten nam o nua phai to giay theo le trinh bay van ban
            # hanh chinh, va cac dong trong khoi can giua VOI NHAU — chuc danh
            # va dong "(Ky va ghi ro ho ten)" nam giua dong ngay thang, khong
            # phai thang hang phai voi no.
            #
            # Cach lam: thut le trai vao dung nua be ngang vung chu roi can
            # giua trong phan con lai. Markdown khong co cu phap nao cho viec
            # nay nen dung mot chi thi rieng, cung kieu voi '\\pagebreak'.
            flush()
            p = doc.add_paragraph()
            body_format(p, spacing=1.3, before=2, after=2,
                        align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.left_indent = Cm(width_cm / 2.0)
            add_inline(p, s[len('\\khoiky '):].strip())
            i += 1; continue
        if s == '\\pagebreak':
            # Dau ngat trang tuong minh. Can cho nhung khoi PHAI nam tron mot
            # trang — vi du khoi ky ten: vo doi giua bang va dong ngay thang
            # thi to giay nhin nhu in hong.
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.page_break_before = True
            p.paragraph_format.space_after = Pt(0)
            i += 1; continue
        if s.startswith('<!--'):
            # Chu thich Markdown la sieu du lieu danh cho bo chuyen (ten de
            # tai, ten ngan tren dau trang), khong phai noi dung de in ra.
            flush()
            while i < len(lines) and '-->' not in lines[i]:
                i += 1
            i += 1; continue
        if s in ('---', '***', '___'):
            flush(); i += 1; continue
        if s.startswith('>'):
            flush(); buf = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                buf.append(lines[i].strip().lstrip('>').strip()); i += 1
            txt = ' '.join(x for x in buf if x)
            if txt:
                p = doc.add_paragraph(); body_format(p, before=8, after=8)
                p.paragraph_format.left_indent = Cm(0.8)
                add_inline(p, txt)
            continue
        m_cap = re.match(r'^\*\*(Bảng|Hình) (\d+)\. (.+)\*\*$', s)
        if m_cap:
            # Chu thich bang dat TREN bang, chu thich hinh dat DUOI hinh —
            # dung quy uoc trinh bay cua khoa. Ca hai cung can giua, in dam,
            # co chu nho hon chu than de khong tranh voi tieu de muc.
            flush()
            cap = doc.add_paragraph()
            body_format(cap, size=12, spacing=1.15, before=8, after=3,
                        align=WD_ALIGN_PARAGRAPH.CENTER)
            set_font(cap.add_run('%s %s. ' % (m_cap.group(1), m_cap.group(2))),
                     12, bold=True)
            set_font(cap.add_run(m_cap.group(3)), 12, bold=True)
            i += 1; continue
        m = re.match(r'^([-*+]|\d+\.)\s+(.*)', s)
        if m:
            flush()
            style = 'List Bullet' if m.group(1) in '-*+' else 'List Number'
            buf = [m.group(2)]; i += 1
            while i < len(lines):
                nxt = lines[i]
                if nxt.strip() and not re.match(r'^([-*+]|\d+\.)\s+', nxt.strip()) \
                   and nxt.startswith(('  ', '\t')):
                    buf.append(nxt.strip()); i += 1
                else:
                    break
            p = doc.add_paragraph(style=style)
            body_format(p, spacing=1.5, before=3, after=3)
            if style == 'List Number':
                if day_so_dang_mo is None:
                    day_so_dang_mo = day_so_moi(doc)
                gan_day_so(p, day_so_dang_mo)
            add_inline(p, ' '.join(buf))
            continue
        para.append(s); i += 1
    flush()

def footer_and_numbering(section, gvhd, svth):
    f = section.footer
    f.is_linked_to_previous = False
    p = f.paragraphs[0]
    p.text = ''
    pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    # hai diem dung tab: giua (SVTH) va phai (so trang); dung API cua python-docx
    pf.tab_stops.add_tab_stop(Cm(8), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT)
    set_font(p.add_run(f'GVHD: {gvhd}'), 11)
    set_font(p.add_run('\t'), 11)
    set_font(p.add_run(f'SVTH: {svth}'), 11)
    set_font(p.add_run('\t'), 11)
    r = p.add_run(); set_font(r, 11)
    for el, attrs, txt in ((('w:fldChar'), {'w:fldCharType': 'begin'}, None),
                           (('w:instrText'), {'xml:space': 'preserve'}, ' PAGE '),
                           (('w:fldChar'), {'w:fldCharType': 'end'}, None)):
        e = OxmlElement(el)
        for k, v in attrs.items(): e.set(qn(k), v)
        if txt: e.text = txt
        r._r.append(e)

LOGO = os.path.join(ROOT, 'docs', 'images', 'logo-tvu.jpg')

def _duong_ke(p, canh='bottom'):
    """Ke mot duong mong o tren hoac duoi doan van — vien cua header/footer."""
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr'); pPr.append(pBdr)
    e = OxmlElement('w:' + canh)
    for k, v in (('w:val', 'single'), ('w:sz', '6'), ('w:space', '4'),
                 ('w:color', '000000')):
        e.set(qn(k), v)
    pBdr.append(e)

def header_de_cuong(section, tieu_de):
    """Dau trang: nhan tai lieu ben trai, ten de tai ben phai, co duong ke duoi."""
    h = section.header
    h.is_linked_to_previous = False
    p = h.paragraphs[0]; p.text = ''
    pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    pf.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT)
    set_font(p.add_run('Đề cương đồ án'), 11, italic=True)
    set_font(p.add_run('\t'), 11)
    set_font(p.add_run(tieu_de.upper()), 11, bold=True, italic=True)
    _duong_ke(p, 'bottom')

TRONG = '. . . . . . . . . . . . . . .'

def footer_de_cuong(section, sv=None):
    """Chan trang: ho ten, MSSV, lop ben trai; so trang ben phai.

    Truong nao chua biet thi de dau cham cho sinh vien dien tay."""
    ten, mssv, lop = (sv or (None, None, None))
    f = section.footer
    f.is_linked_to_previous = False
    p = f.paragraphs[0]; p.text = ''
    pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    pf.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT)
    _duong_ke(p, 'top')
    set_font(p.add_run('%s  -  MSSV %s  -  Lớp %s'
                       % (ten or TRONG, mssv or TRONG, lop or TRONG)), 11)
    set_font(p.add_run('\t'), 11)
    _so_trang(p)

def _so_trang(p):
    """Chen truong PAGE — so trang do Word tu dien, khong go cung."""
    r = p.add_run(); set_font(r, 11)
    for el, attrs, txt in (('w:fldChar', {'w:fldCharType': 'begin'}, None),
                           ('w:instrText', {'xml:space': 'preserve'}, ' PAGE '),
                           ('w:fldChar', {'w:fldCharType': 'end'}, None)):
        e = OxmlElement(el)
        for k, v in attrs.items(): e.set(qn(k), v)
        if txt: e.text = txt
        r._r.append(e)

def trang_bia(doc, tieu_de, loai='ĐỀ CƯƠNG CHI TIẾT',
              noi_ngay=None, sv=None, gvhd=None):
    """Trang bia rieng: khong co dau trang, khong co chan trang, khong co so trang.

    Trang nay dung mot section rieng. Word chi bo dau/chan trang cho section
    dau tien khi section do KHONG lien ket voi section truoc — day cung la ly
    do phai tao section thu hai roi moi dat header/footer len no."""
    from datetime import date
    def dong(txt, co, dam=False, ngh=False, truoc=0, sau=0,
             canh=WD_ALIGN_PARAGRAPH.CENTER):
        p = doc.add_paragraph()
        body_format(p, size=co, spacing=1.2, before=truoc, after=sau, align=canh)
        set_font(p.add_run(txt), co, bold=dam, italic=ngh)
        return p

    dong('TRƯỜNG KỸ THUẬT VÀ CÔNG NGHỆ', 14, truoc=24, sau=2)
    dong('KHOA CÔNG NGHỆ THÔNG TIN', 14, dam=True, sau=18)

    if os.path.exists(LOGO):
        p = doc.add_paragraph()
        body_format(p, before=6, after=18, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.add_run().add_picture(LOGO, width=Cm(3.2))
    else:
        # Khong co logo thi de trong dung phan cho no, khong xe dich bo cuc.
        dong(' ', 13, truoc=40, sau=40)

    dong(loai, 24, dam=True, truoc=12, sau=14)
    dong(tieu_de.upper(), 15, dam=True, sau=48)

    ten, mssv, lop = (sv or (None, None, None))
    for nhan, gt in (('Giảng viên hướng dẫn:',
                      ' ' + (gvhd or '. . . . . . . . . . . . . . . . . . . . . . .')),
                     ('Sinh viên thực hiện:', ''),
                     ('Họ và tên:', ' ' + (ten or '. . . . . . . . . . . . . . . . . . . . . . .')),
                     ('Mã số sinh viên:', ' ' + (mssv or TRONG)),
                     ('Lớp:', ' ' + (lop or TRONG))):
        p = doc.add_paragraph()
        body_format(p, spacing=1.3, before=2, after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.left_indent = Cm(3)
        set_font(p.add_run(nhan), 13, bold=True)
        if gt: set_font(p.add_run(gt), 13)

    # Dong noi - ngay nam sat le duoi trang bia. Word khong co cach neo mot
    # doan vao day trang, nen khoang cach truoc duoc do tu ban xuat: khoi noi
    # dung cua trang bia co do cao co dinh, doan nay ket thuc cach le duoi
    # khoang 1,5 cm. Doi bo cuc trang bia thi phai do lai con so nay.
    dong(noi_ngay or 'Trà Vinh, tháng . . . năm 20 . . .', 13, ngh=True,
         truoc=258)

def new_doc(gvhd='...', svth='...'):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = FONT; st.font.size = Pt(13)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE; pf.line_spacing = 1.5
    pf.space_before = Pt(6); pf.space_after = Pt(6)
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin, s.bottom_margin = Cm(2), Cm(2)
    s.left_margin, s.right_margin = Cm(3), Cm(2)
    footer_and_numbering(s, gvhd, svth)
    return doc, 21 - 3 - 2   # be ngang vung chu, cm

if __name__ == '__main__':
    which = sys.argv[1]           # 'noi-dung' | 'day-du'
    out   = sys.argv[2]
    gvhd  = sys.argv[3] if len(sys.argv) > 3 else '..........................'
    svth  = sys.argv[4] if len(sys.argv) > 4 else '..........................'
    CONTENT = ['01-mo-dau.md','02-chuong-1.md','03-chuong-2.md',
               '04a-chuong-3-thiet-ke.md','04b-chuong-3-cai-dat.md',
               '05-chuong-4.md','06-chuong-5.md']
    SAU = ['07-tai-lieu-tham-khao.md','08a-phu-luc-dac-ta.md','08b-phu-luc-ky-thuat.md']
    # '00-phan-dau.md' KHONG xuat ra Word: tep do la ban dac ta cach trinh bay
    # cac trang bia theo BM5 (kem chu thich co chu, kieu chu), phai dung tay
    # trong Word theo dung bieu mau chu khong phai chep sang.
    # Che do thu tu: mot duong dan tep Markdown bat ky, de xuat de cuong hoac
    # bat cu tai lieu nao khac theo cung mot dinh dang.
    if which.endswith('.md'):
        files = [which if os.path.isabs(which) else os.path.join(ROOT, which)]
    else:
        files = {'noi-dung': CONTENT, 'phu-luc': SAU, 'day-du': CONTENT + SAU}[which]
    # Tep Markdown co dong '<!-- de-cuong: TEN DE TAI -->' o dau duoc xuat
    # theo kieu de cuong: trang bia rieng khong danh so, dau trang mang ten de
    # tai tu trang thu hai tro di.
    ten_de_tai = None
    if len(files) == 1:
        dau_tep = open(files[0], encoding='utf-8').read(400)
        m_dc = re.search(r'<!--\s*de-cuong:\s*(.+?)\s*-->', dau_tep)
        if m_dc:
            ten_de_tai = m_dc.group(1)
        # Ten de tai day du thuong dai qua mot dong o co chu 11pt. Dai qua thi
        # dau trang xuong hai dong va diem dung tab can phai mat tac dung — hai
        # phan dinh lai vao nhau. Nen dau trang dung mot ten ngan rieng.
        m_ng = re.search(r'<!--\s*dau-trang:\s*(.+?)\s*-->', dau_tep)
        ten_dau_trang = m_ng.group(1) if m_ng else ten_de_tai
        # '<!-- sinh-vien: Ho ten | MSSV | Lop -->'. Thieu truong nao thi
        # trang bia va chan trang de dau cham cho cho do.
        m_gv = re.search(r'<!--\s*gvhd:\s*(.+?)\s*-->', dau_tep)
        ten_gvhd = m_gv.group(1) if m_gv else None
        m_nn = re.search(r'<!--\s*noi-ngay:\s*(.+?)\s*-->', dau_tep)
        noi_ngay = m_nn.group(1) if m_nn else None
        m_sv = re.search(r'<!--\s*sinh-vien:\s*(.+?)\s*-->', dau_tep)
        sinh_vien = None
        if m_sv:
            phan = [x.strip() for x in m_sv.group(1).split('|')]
            phan += [''] * (3 - len(phan))
            sinh_vien = tuple(x or None for x in phan[:3])

    if ten_de_tai:
        doc = Document()
        st = doc.styles['Normal']
        st.font.name = FONT; st.font.size = Pt(13)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        pf = st.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE; pf.line_spacing = 1.5
        pf.space_before = Pt(6); pf.space_after = Pt(6)
        for sec in (doc.sections[0],):
            sec.page_width, sec.page_height = Cm(21), Cm(29.7)
            sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
            sec.left_margin, sec.right_margin = Cm(3), Cm(2)
        trang_bia(doc, ten_de_tai, noi_ngay=noi_ngay, sv=sinh_vien,
                  gvhd=ten_gvhd)
        # Section thu hai bat dau o trang moi va mang dau/chan trang; section
        # dau (trang bia) khong dat gi nen no trong — dung y do.
        sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
        sec2.page_width, sec2.page_height = Cm(21), Cm(29.7)
        sec2.top_margin, sec2.bottom_margin = Cm(2), Cm(2)
        sec2.left_margin, sec2.right_margin = Cm(3), Cm(2)
        header_de_cuong(sec2, ten_dau_trang)
        footer_de_cuong(sec2, sinh_vien)
        w = 21 - 3 - 2
    else:
        doc, w = new_doc(gvhd, svth)

    for n, f in enumerate(files):
        path = f if os.path.isabs(f) else os.path.join(DOCS, f)
        render_markdown(doc, path, w, first_chapter_break=(n > 0))
    doc.save(out)
    print(f'  -> {out}  ({len(files)} tệp, {os.path.getsize(out)/1024:.0f} KB)')
    print(f'  hình có chú thích: {len(CAPTIONS)}/{len(FIGS)}')
